"""
ETL lokal: memuat 5 file mentah di data/mentah/ ke database lewat TIGA LEVEL DATA.

    python etl_lokal.py               sinkron sekali (dilewati kalau tidak ada file yang berubah)
    python etl_lokal.py --paksa       muat ulang semuanya walaupun file tidak berubah
    python etl_lokal.py --pantau 5    cek perubahan file setiap 5 menit (sinkronisasi terjadwal)

Alur yang sama seharusnya dikerjakan di Pentaho (lihat SKEMA_DATA.md). Skrip ini dipakai
sebagai cadangan dan pembanding: kalau hasil Pentaho benar, isi tabelnya sama persis.

  LEVEL 1  STAGING     EXTRACT. Isi tiap file disalin apa adanya ke tabel stg_* (semua teks).
                       Ini DUPLIKASI data sumber di dalam database. Hash SHA-256 tiap file dicatat;
                       file yang tidak berubah sejak sinkron terakhir tidak disalin ulang.
  LEVEL 2  INTEGRASI   TRANSFORM + LOAD. Data dibaca dari staging (bukan dari file), lalu
                       diseragamkan: nama & kode provinsi BPS, bentuk tabel, teks jadi angka,
                       buang baris judul/total/Luar Negeri, validasi. Hasil ke tabel bersih.
  LEVEL 3  PRESENTASI  View v_provinsi_lengkap menggabungkan kelima dataset untuk web/API.

Setiap jalan dicatat di etl_batch dan etl_log (SINKRONISASI): kapan, file apa, hash, jumlah baris, status.
"""
import argparse
import csv
import hashlib
import json
import os
import re
import sys
import time
from datetime import datetime

import db

MENTAH = os.path.join(db.BASE_DIR, "data", "mentah")
FILE = {
    "rumah_tangga_listrik": "rumah_tangga_listrik_2024.json",
    "produksi_listrik": "produksi_listrik_2024.csv",
    "kapasitas_pembangkit": "kapasitas_pembangkit_2024.xlsx",
    "sma_listrik": "sma_akses_listrik.xlsx",
    "smk_listrik": "smk_akses_listrik.xlsx",
}
NASIONAL = 99
ALAT = "etl_lokal.py"


def log(msg):
    print(msg, flush=True)


def sekarang():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def hash_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for blok in iter(lambda: f.read(65536), b""):
            h.update(blok)
    return h.hexdigest()


def teks(v):
    return None if v is None or str(v).strip() == "" else str(v)


# ---------------------------------------------------------------- catatan sinkronisasi
def mulai_batch():
    db.execute("INSERT INTO etl_batch (alat, mulai, status) VALUES (?, ?, ?)", (ALAT, sekarang(), "berjalan"))
    return db.query("SELECT MAX(batch_id) AS b FROM etl_batch")[0]["b"]


def selesai_batch(batch, status, catatan=""):
    db.execute("UPDATE etl_batch SET selesai = ?, status = ?, catatan = ? WHERE batch_id = ?",
               (sekarang(), status, catatan[:300], batch))


def catat(batch, tingkat, dataset, status, baris=None, file_asal=None, hash_=None, pesan=""):
    db.execute("INSERT INTO etl_log (batch_id, waktu, tingkat, dataset, file_asal, hash_file, baris, status, pesan) "
               "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
               (batch, sekarang(), tingkat, dataset, file_asal, hash_, baris, status, pesan[:300]))


def hash_terakhir(dataset):
    r = db.query("SELECT hash_file FROM etl_log WHERE tingkat = 'staging' AND dataset = ? AND status = 'ok' "
                 "ORDER BY id DESC LIMIT 1", (dataset,))
    return r[0]["hash_file"] if r else None


# ================================================================ LEVEL 1: STAGING (extract)
def baca_xlsx(path, sheet=None):
    from openpyxl import load_workbook  # pip install openpyxl
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def ekstrak(dataset, path, batch):
    """Baca file sesuai formatnya, kembalikan baris mentah (semua teks) untuk tabel staging."""
    if dataset == "rumah_tangga_listrik":
        # JSON BPS bersarang; diratakan jadi satu baris per (wilayah x klasifikasi), nilainya tetap teks asli
        raw = json.load(open(path, encoding="utf-8"))
        var = str(raw["var"][0]["val"])
        th = raw["tahun"][0]
        turtahun = str(raw["turtahun"][0]["val"])
        rows = []
        for v in raw["vervar"]:
            for t in raw["turvar"]:
                kunci = str(v["val"]) + var + str(t["val"]) + str(th["val"]) + turtahun
                rows.append((batch, str(v["val"]), v["label"], t["label"], str(th["label"]), teks(raw["datacontent"].get(kunci))))
        return ("INSERT INTO stg_rumah_tangga_listrik VALUES (?, ?, ?, ?, ?, ?)", rows)
    if dataset == "produksi_listrik":
        with open(path, encoding="utf-8-sig", newline="") as f:
            data = list(csv.reader(f))
        rows = [(batch, i + 1, teks(r[0] if r else None), teks(r[1] if len(r) > 1 else None)) for i, r in enumerate(data)]
        return ("INSERT INTO stg_produksi_listrik VALUES (?, ?, ?, ?)", rows)
    if dataset == "kapasitas_pembangkit":
        data = baca_xlsx(path)
        rows = [(batch, i + 1, teks(r[0]), teks(r[1] if len(r) > 1 else None)) for i, r in enumerate(data)]
        return ("INSERT INTO stg_kapasitas_pembangkit VALUES (?, ?, ?, ?)", rows)
    # sma / smk: sheet "data", 7 kolom
    data = baca_xlsx(path, "data")
    rows = [(batch, i + 1, *[teks(r[j]) if j < len(r) else None for j in range(7)]) for i, r in enumerate(data)]
    return (f"INSERT INTO stg_{dataset} VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", rows)


def level_staging(batch, paksa):
    """Salin file yang berubah ke staging. Kembalikan True kalau ada yang berubah."""
    log("LEVEL 1  STAGING (extract, duplikasi data sumber)")
    berubah = False
    for dataset, nama_file in FILE.items():
        path = os.path.join(MENTAH, nama_file)
        h = hash_file(path)
        if not paksa and h == hash_terakhir(dataset):
            catat(batch, "staging", dataset, "tidak berubah", file_asal=nama_file, hash_=h,
                  pesan="Hash sama dengan sinkron terakhir, staging dipakai ulang")
            log(f"  {dataset:<22} tidak berubah ({h[:10]}…)")
            continue
        sql, rows = ekstrak(dataset, path, batch)
        db.execute(f"DELETE FROM stg_{dataset}")
        db.executemany(sql, rows)
        catat(batch, "staging", dataset, "ok", len(rows), nama_file, h, f"Disalin dari {nama_file}")
        log(f"  {dataset:<22} {len(rows):>4} baris mentah  <- {nama_file}  ({h[:10]}…)")
        berubah = True
    return berubah


# ================================================================ LEVEL 2: INTEGRASI (transform + load)
def normal(nama):
    """Kunci pencocokan nama provinsi: huruf besar, 'KEP.' disamakan, spasi dirapikan."""
    s = str(nama).replace("﻿", "").strip().upper()
    s = re.sub(r"\s+", " ", s)
    s = s.replace("KEPULAUAN ", "KEP. ")
    s = s.replace("DAERAH ISTIMEWA YOGYAKARTA", "DI YOGYAKARTA")
    return s


def nama_tampil(label):
    khusus = {"DKI JAKARTA": "DKI Jakarta", "DI YOGYAKARTA": "DI Yogyakarta", "INDONESIA": "Indonesia"}
    if label in khusus:
        return khusus[label]
    label = label.replace("KEP. ", "KEPULAUAN ")
    return " ".join(w.capitalize() for w in label.lower().split())


def wilayah(kode):
    if kode == NASIONAL:
        return "Nasional"
    for batas, nama in ((30, "Sumatera"), (40, "Jawa"), (60, "Bali dan Nusa Tenggara"),
                        (70, "Kalimantan"), (80, "Sulawesi"), (90, "Maluku"), (99, "Papua")):
        if kode < batas:
            return nama
    return "Lainnya"


def angka(v):
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "-", "…", "..."):
        return None
    return float(s.replace(",", "."))


def bilangan_bulat(v):
    s = "" if v is None else str(v).strip()
    return int(float(s)) if re.fullmatch(r"\d+(\.0+)?", s) else None


def t_rumah_tangga():
    """Pivot: satu baris per wilayah x klasifikasi -> satu baris per provinsi (kota, desa, total)."""
    rows = db.query("SELECT kode_wilayah, nama_wilayah, klasifikasi, tahun, nilai FROM stg_rumah_tangga_listrik")
    per = {}
    tahun = None
    for r in rows:
        kode = int(r["kode_wilayah"]) // 100           # kode 4 digit BPS (1100) -> kode provinsi (11)
        tahun = int(r["tahun"])
        d = per.setdefault(kode, {"nama": nama_tampil(normal(r["nama_wilayah"])), "nilai": {}})
        d["nilai"][r["klasifikasi"]] = angka(r["nilai"])
    provinsi, baris = [], []
    for kode in sorted(per):
        d = per[kode]
        provinsi.append((kode, d["nama"], wilayah(kode)))
        n = d["nilai"]
        baris.append((kode, tahun, n.get("Perkotaan"), n.get("Perdesaan"), n.get("Perkotaan+Perdesaan")))
    return provinsi, baris, tahun


def t_dua_kolom(tabel, peta_nama):
    """CSV/XLSX ekspor BPS: kolom 1 nama provinsi, kolom 2 nilai; baris judul & tahun di atas."""
    rows = db.query(f"SELECT kolom_1, kolom_2 FROM {tabel} ORDER BY baris_ke")
    tahun, hasil, tak_cocok = None, [], []
    for r in rows:
        a, b = r["kolom_1"], r["kolom_2"]
        if a is None:
            if b and re.fullmatch(r"\d{4}", b.strip()):
                tahun = int(b.strip())
            continue
        kode = peta_nama.get(normal(a))
        nilai = angka(b)
        if kode is None:
            if nilai is not None:
                tak_cocok.append(a)
            continue
        hasil.append((kode, nilai))
    return hasil, tahun, tak_cocok


def t_sekolah(tabel, kode_valid):
    rows = db.query(f"SELECT * FROM {tabel} ORDER BY baris_ke")
    judul = rows[0]["kolom_1"] or ""
    m = re.search(r"Tahun\s+(\d{4})", judul)
    tahun = int(m.group(1)) if m else None          # tahun dari judul, bukan dari nama file
    hasil = []
    for r in rows:
        if normal(r["kolom_4"] or "") == "INDONESIA":
            kode = NASIONAL
        else:
            kode = bilangan_bulat(r["kolom_3"])     # Kode BPS, bukan Kode Kemdagri (beda untuk Papua)
            if kode is None:
                continue                            # baris judul & "Luar Negeri" (tanpa kode) dibuang
        if kode not in kode_valid:
            continue
        hasil.append((kode, tahun, angka(r["kolom_5"]), angka(r["kolom_6"]), angka(r["kolom_7"])))
    return hasil, tahun


def cek(kondisi, pesan, masalah):
    if not kondisi:
        masalah.append(pesan)


def level_integrasi(batch):
    log("LEVEL 2  INTEGRASI (transform dari staging, validasi, load)")
    masalah = []
    prov, rt, tahun_energi = t_rumah_tangga()
    kode_valid = {p[0] for p in prov}
    peta_nama = {normal(p[1]): p[0] for p in prov}
    produksi, th_csv, tak1 = t_dua_kolom("stg_produksi_listrik", peta_nama)
    kapasitas, th_kap, tak2 = t_dua_kolom("stg_kapasitas_pembangkit", peta_nama)
    sma, th_sma = t_sekolah("stg_sma_listrik", kode_valid)
    smk, th_smk = t_sekolah("stg_smk_listrik", kode_valid)
    log(f"  transform: rumah tangga {len(rt)}, produksi {len(produksi)}, kapasitas {len(kapasitas)}, "
        f"SMA {len(sma)} (th {th_sma}), SMK {len(smk)} (th {th_smk})")
    kosong = [p[1] for p, r in zip(prov, rt) if None in r[2:]]
    if kosong:
        log(f"  nilai kosong dibiarkan NULL, bukan 0: {', '.join(kosong)}")

    # validasi
    for nama, rows in (("rumah tangga", rt), ("produksi", produksi), ("kapasitas", kapasitas), ("SMA", sma), ("SMK", smk)):
        kode_ada = {r[0] for r in rows}
        cek(kode_ada == kode_valid, f"{nama}: provinsi tidak lengkap, kurang {sorted(kode_valid - kode_ada)}", masalah)
    for nama, tak in (("produksi", tak1), ("kapasitas", tak2)):
        cek(not tak, f"{nama}: nama provinsi tidak dikenali: {tak}", masalah)
    for nama, rows in (("produksi", produksi), ("kapasitas", kapasitas)):
        nas = next((v for k, v in rows if k == NASIONAL), None)
        jumlah = sum(v for k, v in rows if k != NASIONAL and v is not None)
        cek(nas is not None and abs(jumlah - nas) < 0.1,
            f"{nama}: jumlah 38 provinsi {jumlah:.2f} tidak sama dengan baris INDONESIA {nas}", masalah)
        log(f"  cek {nama:<10} jumlah 38 provinsi = {jumlah:,.2f}, baris INDONESIA = {nas:,.2f}")
    for nama, rows in (("rumah tangga", rt), ("SMA", sma), ("SMK", smk)):
        for r in rows:
            for v in r[2:]:
                cek(v is None or 0 <= v <= 100, f"{nama}: persentase di luar 0-100 pada kode {r[0]}", masalah)
    if masalah:
        for m in masalah:
            log("  GAGAL: " + m)
            catat(batch, "integrasi", "validasi", "gagal", pesan=m)
        return False
    log("  semua cek lolos")

    muat = [
        ("provinsi", "INSERT INTO provinsi (kode_bps, nama, wilayah) VALUES (?, ?, ?)", prov),
        ("rumah_tangga_listrik", "INSERT INTO rumah_tangga_listrik VALUES (?, ?, ?, ?, ?)", rt),
        ("produksi_listrik", "INSERT INTO produksi_listrik VALUES (?, ?, ?)", [(k, th_csv, v) for k, v in produksi]),
        ("kapasitas_pembangkit", "INSERT INTO kapasitas_pembangkit VALUES (?, ?, ?)", [(k, th_kap, v) for k, v in kapasitas]),
        ("sma_listrik", "INSERT INTO sma_listrik VALUES (?, ?, ?, ?, ?)", sma),
        ("smk_listrik", "INSERT INTO smk_listrik VALUES (?, ?, ?, ?, ?)", smk),
    ]
    for tabel, sql, rows in muat:
        db.execute(f"DELETE FROM {tabel}")
        db.executemany(sql, rows)
        catat(batch, "integrasi", tabel, "ok", len(rows), pesan="Dimuat ulang dari staging")
        log(f"  {tabel:<22} {len(rows)} baris")

    url_bps = "https://www.bps.go.id/id/statistics-table?subject=558"
    sumber = [
        ("rumah_tangga_listrik", "Persentase rumah tangga dengan sumber penerangan utama dari listrik, menurut provinsi dan klasifikasi desa",
         "JSON", "BPS, Susenas Maret", url_bps, "Persen rumah tangga", FILE["rumah_tangga_listrik"], tahun_energi),
        ("produksi_listrik", "Tenaga listrik yang dibangkitkan menurut provinsi",
         "CSV", "Badan Pusat Statistik (BPS)", "https://www.bps.go.id", "GWh", FILE["produksi_listrik"], th_csv),
        ("kapasitas_pembangkit", "Kapasitas terpasang pembangkit listrik menurut provinsi",
         "XLSX", "Badan Pusat Statistik (BPS)", "https://www.bps.go.id", "MW", FILE["kapasitas_pembangkit"], th_kap),
        ("sma_listrik", "Proporsi SMA yang memiliki akses listrik",
         "XLSX", "Satu Data, Pusdatin Kementerian Pendidikan Dasar dan Menengah", "", "Persen satuan pendidikan", FILE["sma_listrik"], th_sma),
        ("smk_listrik", "Proporsi SMK yang memiliki akses listrik",
         "XLSX", "Satu Data, Pusdatin Kementerian Pendidikan Dasar dan Menengah", "", "Persen satuan pendidikan", FILE["smk_listrik"], th_smk),
    ]
    db.execute("DELETE FROM sumber_data")
    db.executemany("INSERT INTO sumber_data (kode, nama, format_asal, sumber, url, keterangan, file_asal, tahun_data, data_contoh) "
                   "VALUES (?, ?, ?, ?, ?, ?, ?, ?, 0)", sumber)
    log(f"  {'sumber_data':<22} {len(sumber)} baris")
    return True


# ================================================================ LEVEL 3: PRESENTASI
def level_presentasi(batch):
    log("LEVEL 3  PRESENTASI (view gabungan untuk web)")
    n = db.query("SELECT COUNT(*) AS n FROM v_provinsi_lengkap")[0]["n"]
    lengkap = db.query("SELECT COUNT(*) AS n FROM v_provinsi_lengkap WHERE produksi_gwh IS NOT NULL AND kapasitas_mw IS NOT NULL "
                       "AND rumah_tangga_persen IS NOT NULL AND sma_persen IS NOT NULL AND smk_persen IS NOT NULL")[0]["n"]
    catat(batch, "presentasi", "v_provinsi_lengkap", "ok", n, pesan=f"{lengkap} dari {n} baris punya kelima dataset")
    log(f"  v_provinsi_lengkap     {n} baris, {lengkap} lengkap kelima dataset")


# ================================================================ jalankan
def sinkron(paksa=False):
    if db.mode() == "dw":
        log("Database ini memakai data warehouse hasil ETL Pentaho (schema dw).")
        log("etl_lokal.py tidak dijalankan supaya data Pentaho tidak tertimpa.")
        log("Kalau memang mau memakai ETL lokal, set SUMBER_DATA=lokal di .env.")
        return True
    db.init_schema()
    batch = mulai_batch()
    log(f"== Sinkronisasi batch {batch}, {sekarang()} ==")
    try:
        if not level_staging(batch, paksa):
            selesai_batch(batch, "sudah sinkron", "Tidak ada file sumber yang berubah")
            log("Tidak ada file yang berubah. Database sudah sinkron.")
            return True
        if not level_integrasi(batch):
            selesai_batch(batch, "gagal", "Validasi gagal, tabel integrasi tidak diubah")
            log("GAGAL. Tabel integrasi tidak diubah.")
            return False
        level_presentasi(batch)
        selesai_batch(batch, "berhasil")
        log("Selesai.")
        return True
    except Exception as e:
        selesai_batch(batch, "gagal", f"{type(e).__name__}: {e}")
        raise


def main():
    ap = argparse.ArgumentParser(description="ETL 3 level: file -> staging -> integrasi -> presentasi")
    ap.add_argument("--paksa", action="store_true", help="muat ulang walaupun file tidak berubah")
    ap.add_argument("--pantau", type=float, metavar="MENIT", help="ulangi pengecekan setiap N menit")
    a = ap.parse_args()
    if a.pantau:
        log(f"Mode pantau: cek perubahan file setiap {a.pantau:g} menit. Ctrl+C untuk berhenti.")
        while True:
            sinkron(a.paksa)
            a.paksa = False
            time.sleep(a.pantau * 60)
    ok = sinkron(a.paksa)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
